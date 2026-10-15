"""Colours the client's rider / driver spec files by what is built: green = done, red = left.

Run `python docs/modules/Bookings/make_status.py` after a batch of work. An item is done when it is
in RIDER_DONE / DRIVER_DONE below (built before `docs/rides_tasks.md` existed) or ticked `- [x]` in
`docs/rides_tasks.md` — the task list the team works from. Writes `*_status.docx` next to the
originals, which stay untouched.
"""
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Pt

BASE = str(Path(__file__).resolve().parent)
TASKS = Path(__file__).resolve().parents[2] / "rides_tasks.md"


def nums(spec):
    out = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-")
            out.update(str(i) for i in range(int(a), int(b) + 1))
        else:
            out.add(part)
    return out


RIDER_DONE = nums(
    "1-4,7,8,10,11,15,19,23,30,31,43,44,46,50,51,54-60,62-74,76,78-80,82,83,85,87-90,96-104,109,122,"
    "124-126,128-132,141-143,145,163,164,166-169,173,175-178,185,186,190-198,200,202-207,210-215,217,"
    "249-251,"
    # phase 3 — booking ahead
    "12,105-108,165,170,"
    # safety + changing the trip
    "86,91-95,110,111"
) | {"R-NEW-01", "R-NEW-12"}

DRIVER_DONE = nums(
    "1-4,7,9,11-15,21,24-27,40,45,46,48-58,62-70,74,81,85-87,103-110,113,115-118,120,121,123,125-128,"
    "135-139,161-167,171-174,176,178-181,183-186,189-192,198,199,203,206-209,225,255,305,306,315,"
    "319-323,325-328,331-334,342,343,345,346-358,360-364,366,367,371,372,386,387,400,"
    # phase 3 — booking ahead
    "119,168-170,260-269,385,"
    # safety + changing the trip
    "131-134,141,142,156,383"
) | {"D-NEW-01", "D-NEW-12"}



def ticked():
    """Items ticked `- [x] **ID**` in docs/rides_tasks.md: R-12 / R-INN-01 → rider, D-12 / D-NEW-01 → driver, AI-xx → both."""
    rider, driver = set(), set()
    if not TASKS.exists():
        return rider, driver
    for raw in TASKS.read_text(encoding="utf-8").splitlines():
        m = re.match(r"-\s*\[[xX]\]\s*\*\*([A-Z]+-[A-Z0-9-]+)\*\*", raw.strip())
        if not m:
            continue
        tid = m.group(1)
        if tid.startswith("AI-"):
            rider.add(tid)
            driver.add(tid)
        elif re.fullmatch(r"R-\d+", tid):
            rider.add(tid[2:])
        elif re.fullmatch(r"D-\d+", tid):
            driver.add(tid[2:])
        elif tid.startswith("R-"):
            rider.add(tid)
        elif tid.startswith("D-"):
            driver.add(tid)
    return rider, driver


GREEN_BG, RED_BG = "C6EFCE", "FFC7CE"
GREEN_TX, RED_TX = RGBColor(0x00, 0x61, 0x00), RGBColor(0x9C, 0x00, 0x06)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def is_item(text):
    t = text.strip()
    return t.isdigit() or t.startswith(("R-", "D-", "AI-"))


def run(src, dst, done, title):
    doc = Document(src)
    counts = {"done": 0, "left": 0}
    seen = set()
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if not cells:
                continue
            key = cells[0].text.strip()
            if not is_item(key):
                continue
            ok = key in done
            counts["done" if ok else "left"] += 1
            seen.add(key)
            done_cells = set()
            for c in cells:
                if id(c._tc) in done_cells:
                    continue
                done_cells.add(id(c._tc))
                shade(c, GREEN_BG if ok else RED_BG)
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = GREEN_TX if ok else RED_TX
    missing = done - seen
    total = counts["done"] + counts["left"]

    body = doc.element.body
    first = body[0]
    lines = [
        (f"حالة التنفيذ — {title} (تحديث {date.today().isoformat()})", True, None),
        (f"🟩 أخضر = تم تنفيذه في الكود: {counts['done']} بند", True, GREEN_TX),
        (f"🟥 أحمر = لم يُنفَّذ بعد: {counts['left']} بند", True, RED_TX),
        (f"الإجمالي: {total} بند — نسبة الإنجاز {round(100 * counts['done'] / total)}%", False, None),
        ("البند الجزئي يُحسب أحمر حتى يكتمل. المتبقي ومين شغال على إيه: docs/rides_tasks.md", False, None),
    ]
    for text, bold, color in lines:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = 2  # right
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(12)
        rPr = r._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)
        if color is not None:
            r.font.color.rgb = color
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        pPr.insert(0, bidi)
        body.remove(p._p)
        first.addprevious(p._p)
    doc.save(dst)
    print(title, counts, "total", total, "unmatched-done:", sorted(missing))


if __name__ == "__main__":
    tick_rider, tick_driver = ticked()
    RIDER_DONE |= tick_rider
    DRIVER_DONE |= tick_driver
    run(BASE + r"\DORR_Rider_Merged_All_Features_and_Ideas.docx",
        BASE + r"\DORR_Rider_Merged_All_Features_and_Ideas_status.docx", RIDER_DONE, "تطبيق الراكب")
    run(BASE + r"\DORR_Driver_Merged_All_Features_and_Ideas.docx",
        BASE + r"\DORR_Driver_Merged_All_Features_and_Ideas_status.docx", DRIVER_DONE, "تطبيق السائق")
